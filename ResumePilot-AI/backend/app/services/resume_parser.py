import io
import re

import fitz
from docx import Document


SKILLS = [
    "python", "java", "javascript", "typescript", "c++", "c", "sql",
    "html", "css", "react", "node.js", "node", "express", "fastapi", "flask",
    "spring boot", "spring", "mongodb", "mysql", "postgresql", "docker",
    "kubernetes", "git", "github", "aws", "azure", "gcp", "tensorflow",
    "pytorch", "scikit-learn", "pandas", "numpy", "opencv", "langchain",
    "rag", "llm", "machine learning", "deep learning", "rest api",
    "microservices", "linux", "data structures", "algorithms", "oops",
    "object oriented programming", "communication", "problem solving"
]

ALIASES = {
    "node": "node.js",
    "spring": "spring boot",
    "object oriented programming": "oops"
}


def extract_resume(filename: str, content: bytes) -> str:
    filename = filename.lower()

    if filename.endswith(".pdf"):
        return extract_pdf(content)
    if filename.endswith(".docx"):
        return extract_docx(content)
    if filename.endswith(".txt"):
        return content.decode("utf-8", errors="ignore")

    raise ValueError(
        "Unsupported file format. Please upload PDF, DOCX or TXT."
    )


def extract_pdf(content: bytes) -> str:
    text = []

    with fitz.open(stream=content, filetype="pdf") as pdf:
        for page in pdf:
            page_text = page.get_text()
            if page_text:
                text.append(page_text)

    return "\n".join(text).strip()


def extract_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())

            if row_text:
                text.append(" | ".join(row_text))

    return "\n".join(text).strip()


def normalize_text(text: str) -> str:
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_skills(text: str) -> list:
    normalized_text = normalize_text(text)
    found_skills = []

    for skill in SKILLS:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"

        if re.search(pattern, normalized_text):
            canonical_skill = ALIASES.get(skill, skill)

            if canonical_skill not in found_skills:
                found_skills.append(canonical_skill)

    return found_skills


def extract_sections(text: str) -> dict:
    lines = text.splitlines()

    sections = {
        "education": [],
        "skills": [],
        "experience": [],
        "projects": [],
        "certifications": [],
        "achievements": []
    }

    current_section = None

    section_mapping = {
        "education": "education",
        "academic": "education",
        "skills": "skills",
        "technical skills": "skills",
        "experience": "experience",
        "work experience": "experience",
        "internship": "experience",
        "projects": "projects",
        "project": "projects",
        "certifications": "certifications",
        "certificates": "certifications",
        "achievements": "achievements",
        "achievement": "achievements"
    }

    for line in lines:
        clean_line = line.strip()

        if not clean_line:
            continue

        normalized_line = clean_line.lower().replace(":", "")

        if normalized_line in section_mapping:
            current_section = section_mapping[normalized_line]
            continue

        if current_section:
            sections[current_section].append(clean_line)

    return sections


def parse_resume(filename: str, content: bytes) -> dict:
    raw_text = extract_resume(filename, content)

    if not raw_text:
        raise ValueError("Could not extract text from resume.")

    return {
        "filename": filename,
        "raw_text": raw_text,
        "skills": extract_skills(raw_text),
        "sections": extract_sections(raw_text)
    }
